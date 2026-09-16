# -*- coding: utf-8 -*-
"""VBA から見た Excel オブジェクトモデルのモック。

SaveAs は openpyxl で本物の .xlsx を書き、Workbooks.Open はそれを読み戻す。
つまり Verify は「保存されたファイル」を検査しており、メモリ上の模型を
見ているわけではない。DisplayFormat は条件付き書式を評価して、Excel が
実際に表示する書式を返す。
"""
import datetime, re, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, column_index_from_string

from vbaint import VDate, EMPTY, NOTHING, MISSING, VBAError, vba_str, num, to_date, EPOCH

MAXROW, MAXCOL = 1048576, 16384

# -------------------------------------------------------------- colours

def bgr_to_hex(c):
    c = int(c)
    r, g, b = c & 255, (c >> 8) & 255, (c >> 16) & 255
    return "%02X%02X%02X" % (r, g, b)

def hex_to_bgr(h):
    if h is None: return None
    h = str(h)
    if len(h) == 8: h = h[2:]
    if len(h) != 6: return None
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return r + g * 256 + b * 65536

# VBA の罫線定数 <-> openpyxl のスタイル名
LS_CONT, LS_DASH, LS_NONE = 1, -4115, -4142
W_HAIR, W_THIN, W_MEDIUM, W_THICK = 1, 2, -4138, 4

def side_name(linestyle, weight):
    if linestyle == LS_NONE or linestyle is None: return None
    if linestyle == LS_DASH:
        return {W_THICK: "mediumDashed", W_MEDIUM: "mediumDashed"}.get(weight, "dashed")
    return {W_HAIR: "hair", W_THIN: "thin", W_MEDIUM: "medium", W_THICK: "thick"}.get(weight, "thin")

def name_to_ls_weight(name):
    if name is None: return (LS_NONE, W_THIN)
    if name in ("dashed", "mediumDashed", "dashDot", "dashDotDot"):
        return (LS_DASH, W_MEDIUM if name == "mediumDashed" else W_THIN)
    return (LS_CONT, {"hair": W_HAIR, "thin": W_THIN, "medium": W_MEDIUM,
                      "thick": W_THICK}.get(name, W_THIN))

EDGE_LEFT, EDGE_TOP, EDGE_BOTTOM, EDGE_RIGHT = 7, 8, 9, 10
FC_LEFT, FC_RIGHT, FC_TOP, FC_BOTTOM = -4131, -4152, -4160, -4107
EDGE_KEY = {EDGE_LEFT: "left", EDGE_TOP: "top", EDGE_BOTTOM: "bottom", EDGE_RIGHT: "right",
            FC_LEFT: "left", FC_TOP: "top", FC_BOTTOM: "bottom", FC_RIGHT: "right"}

ALIGN = {-4131: "left", -4152: "right", -4108: "center", 1: "general"}

# -------------------------------------------------------------- data model

class CellData:
    __slots__ = ("value", "formula", "number_format", "bold", "font_color",
                 "fill", "borders", "halign", "wrap", "locked")
    def __init__(self):
        self.value = None
        self.formula = None
        self.number_format = "General"
        self.bold = False
        self.font_color = None
        self.fill = None
        self.borders = {}
        self.halign = None
        self.wrap = None
        self.locked = None
    def clone(self):
        c = CellData()
        for s in CellData.__slots__:
            v = getattr(self, s)
            setattr(c, s, dict(v) if isinstance(v, dict) else v)
        return c

class CFRule:
    __slots__ = ("r1", "c1", "r2", "c2", "formula", "fill", "borders")
    def __init__(self, r1, c1, r2, c2, formula):
        self.r1, self.c1, self.r2, self.c2 = r1, c1, r2, c2
        self.formula = formula
        self.fill = None
        self.borders = {}

class SheetData:
    def __init__(self, name):
        self.name = name
        self.cells = {}
        self.col_width = {}
        self.col_hidden = set()
        self.merges = []
        self.cf = []
        self.validations = {}
        self.protected = False
        self.visible = "visible"
        self.freeze = None
        self.default_locked = True
        self.shapes = []
    def cell(self, r, c, create=True):
        cd = self.cells.get((r, c))
        if cd is None and create:
            cd = CellData(); self.cells[(r, c)] = cd
        return cd
    def locked_at(self, r, c):
        cd = self.cells.get((r, c))
        if cd is None or cd.locked is None: return self.default_locked
        return cd.locked
    def clone(self, name=None):
        s = SheetData(name or self.name)
        s.cells = {k: v.clone() for k, v in self.cells.items()}
        s.col_width = dict(self.col_width)
        s.col_hidden = set(self.col_hidden)
        s.merges = list(self.merges)
        s.cf = [copy.copy(r) for r in self.cf]
        for r in s.cf: r.borders = dict(r.borders)
        s.validations = dict(self.validations)
        s.protected = self.protected
        s.visible = self.visible
        s.freeze = self.freeze
        s.default_locked = self.default_locked
        return s
    def max_used(self):
        if not self.cells: return (0, 0)
        rs = [k[0] for k in self.cells]; cs = [k[1] for k in self.cells]
        return (max(rs), max(cs))

# -------------------------------------------------------------- CF formula

TOKEN = re.compile(r'\s*(?:(?P<str>"[^"]*")|(?P<fn>[A-Z]+)\s*\(|(?P<ref>\$?[A-Z]{1,3}\$?\d+)'
                   r'|(?P<num>\d+(?:\.\d+)?)|(?P<op><=|>=|<>|[<>=])|(?P<p>[(),]))')

class FormulaEval:
    """条件付き書式の式を評価する小さな評価器。

    行 2 の日付と、その行の $C/$D を参照する式だけを想定しているが、
    参照の相対/絶対とオフセット計算は一般的に実装してある。
    """
    def __init__(self, sheet, anchor_r, anchor_c, cell_r, cell_c):
        self.sh = sheet
        self.dr, self.dc = cell_r - anchor_r, cell_c - anchor_c

    def run(self, formula):
        s = formula[1:] if formula.startswith("=") else formula
        self.s, self.i = s, 0
        v = self.expr()
        return bool(v)

    # -- lexing
    def peek(self):
        m = TOKEN.match(self.s, self.i)
        return m
    def take(self):
        m = TOKEN.match(self.s, self.i)
        if not m: raise ValueError("bad formula at %r" % self.s[self.i:])
        self.i = m.end()
        return m

    def expr(self):
        left = self.operand()
        m = self.peek()
        if m and m.lastgroup == "op":
            self.take()
            op = m.group("op")
            right = self.operand()
            return self.compare(op, left, right)
        return left

    def compare(self, op, a, b):
        if a is None: a = ""
        if b is None: b = ""
        if isinstance(a, str) or isinstance(b, str):
            if not (isinstance(a, str) and isinstance(b, str)):
                # 片方だけ文字列: Excel では数値 < 文字列
                if op == "=": return False
                if op == "<>": return True
                return isinstance(b, str)
            r = (a > b) - (a < b)
        else:
            x, y = float(a), float(b)
            r = (x > y) - (x < y)
        return {"=": r == 0, "<>": r != 0, "<": r < 0, ">": r > 0,
                "<=": r <= 0, ">=": r >= 0}[op]

    def operand(self):
        m = self.take()
        g = m.lastgroup
        if g == "str": return m.group("str")[1:-1]
        if g == "num": return float(m.group("num"))
        if g == "ref": return self.ref_value(m.group("ref"))
        if g == "fn":
            fn = m.group("fn")
            args = []
            while True:
                args.append(self.expr())
                m2 = self.take()
                if m2.group("p") == ")": break
                if m2.group("p") != ",": raise ValueError("expected , or )")
            if fn == "AND": return all(bool(a) for a in args)
            if fn == "OR": return any(bool(a) for a in args)
            if fn == "NOT": return not bool(args[0])
            if fn == "WEEKDAY":
                d = float(args[0]); base = int(args[1]) if len(args) > 1 else 1
                dow = (int(d) + 6) % 7 + 1
                return float((dow - base) % 7 + 1) if base != 2 else float((dow - 2) % 7 + 1)
            if fn == "TRUE": return True
            if fn == "FALSE": return False
            raise ValueError("unsupported function " + fn)
        if g == "p" and m.group("p") == "(":
            v = self.expr(); self.take()
            return v
        raise ValueError("unexpected token")

    def ref_value(self, ref):
        m = re.match(r"(\$?)([A-Z]{1,3})(\$?)(\d+)", ref)
        dollar_c, col, dollar_r, row = m.groups()
        c = column_index_from_string(col) + (0 if dollar_c else self.dc)
        r = int(row) + (0 if dollar_r else self.dr)
        cd = self.sh.cells.get((r, c))
        if cd is None or cd.value is None: return ""
        v = cd.value
        if isinstance(v, VDate): return float(v)
        if isinstance(v, bool): return 1.0 if v else 0.0
        return v

def effective_format(sheet, r, c):
    """条件付き書式を適用したあとの塗り / 罫線 (= Excel が表示するもの)。"""
    cd = sheet.cells.get((r, c))
    fill = cd.fill if cd else None
    borders = dict(cd.borders) if cd else {}
    fill_set = False
    border_set = set()
    for rule in sheet.cf:
        if not (rule.r1 <= r <= rule.r2 and rule.c1 <= c <= rule.c2): continue
        try:
            hit = FormulaEval(sheet, rule.r1, rule.c1, r, c).run(rule.formula)
        except Exception:
            hit = False
        if not hit: continue
        if rule.fill is not None and not fill_set:
            fill = rule.fill; fill_set = True
        for k, v in rule.borders.items():
            if k not in border_set:
                borders[k] = v; border_set.add(k)
    return fill, borders

# -------------------------------------------------------------- xlsx I/O

def _to_py(v):
    if isinstance(v, VDate):
        d = EPOCH + datetime.timedelta(days=int(v))
        frac = float(v) - int(v)
        if frac == 0: return datetime.datetime(d.year, d.month, d.day)
        return datetime.datetime(d.year, d.month, d.day) + datetime.timedelta(seconds=round(frac*86400))
    if v is EMPTY or v is None: return None
    return v

def _from_py(v):
    if isinstance(v, datetime.datetime):
        return VDate((v.date() - EPOCH).days + (v.hour*3600+v.minute*60+v.second)/86400.0)
    if isinstance(v, datetime.date):
        return VDate((v - EPOCH).days)
    return v

def save_workbook(wbd, path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sd in wbd.sheets:
        ws = wb.create_sheet(sd.name)
        ws.sheet_state = sd.visible
        for (r, c), cd in sorted(sd.cells.items()):
            cell = ws.cell(row=r, column=c)
            if cd.formula is not None:
                cell.value = cd.formula
            else:
                cell.value = _to_py(cd.value)
            if cd.number_format and cd.number_format != "General":
                cell.number_format = cd.number_format
            if cd.bold or cd.font_color is not None:
                cell.font = Font(bold=bool(cd.bold),
                                 color=(bgr_to_hex(cd.font_color) if cd.font_color is not None else None))
            if cd.fill is not None:
                h = bgr_to_hex(cd.fill)
                cell.fill = PatternFill(fill_type="solid", start_color=h, end_color=h)
            if cd.borders:
                kw = {}
                for k, (ls, col, wt) in cd.borders.items():
                    nm = side_name(ls, wt)
                    if nm: kw[k] = Side(style=nm, color=bgr_to_hex(col) if col is not None else None)
                if kw: cell.border = Border(**kw)
            if cd.halign or cd.wrap is not None:
                cell.alignment = Alignment(horizontal=cd.halign, wrap_text=bool(cd.wrap))
            lk = sd.locked_at(r, c)
            cell.protection = Protection(locked=bool(lk))
        for c, w in sd.col_width.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        for c in sd.col_hidden:
            ws.column_dimensions[get_column_letter(c)].hidden = True
        for (r1, c1, r2, c2) in sd.merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for rule in sd.cf:
            ref = "%s%d:%s%d" % (get_column_letter(rule.c1), rule.r1,
                                 get_column_letter(rule.c2), rule.r2)
            fill = None
            if rule.fill is not None:
                h = bgr_to_hex(rule.fill)
                fill = PatternFill(patternType="solid", start_color=h, end_color=h)
            border = None
            if rule.borders:
                kw = {}
                for k, (ls, col, wt) in rule.borders.items():
                    nm = side_name(ls, wt)
                    if nm: kw[k] = Side(style=nm, color=bgr_to_hex(col) if col is not None else None)
                if kw: border = Border(**kw)
            dxf = DifferentialStyle(fill=fill, border=border)
            ws.conditional_formatting.add(ref, Rule(type="expression", dxf=dxf,
                                                    formula=[rule.formula.lstrip("=")]))
        for (r, c), vd in sd.validations.items():
            addr = "%s%d" % (get_column_letter(c), r)
            dv = DataValidation(type=vd["type"], operator=vd["op"],
                                formula1=vd["f1"], formula2=vd["f2"],
                                allow_blank=False, showErrorMessage=True,
                                errorTitle=vd.get("title"), error=vd.get("msg"))
            ws.add_data_validation(dv)
            dv.add(addr)
        if sd.protected:
            ws.protection.sheet = True
            ws.protection.insertRows = False
            ws.protection.deleteRows = False
            ws.protection.insertColumns = False
            ws.protection.deleteColumns = False
            ws.protection.sort = False
            ws.protection.autoFilter = False
        if sd.freeze:
            ws.freeze_panes = "%s%d" % (get_column_letter(sd.freeze[1] + 1), sd.freeze[0] + 1)
    if wbd.protect_structure:
        wb.security = openpyxl.workbook.protection.WorkbookProtection(lockStructure=True)
    wb.save(path)

def load_workbook_data(path):
    wb = openpyxl.load_workbook(path)
    sheets = []
    for ws in wb.worksheets:
        sd = SheetData(ws.title)
        sd.visible = ws.sheet_state
        sd.protected = bool(ws.protection.sheet)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and cell.has_style is False: continue
                cd = CellData()
                v = _from_py(cell.value)
                if isinstance(v, str) and v.startswith("="):
                    cd.formula = v; cd.value = None
                else:
                    cd.value = v
                cd.number_format = cell.number_format
                f = cell.font
                cd.bold = bool(f.bold)
                if f.color is not None and getattr(f.color, "rgb", None):
                    cd.font_color = hex_to_bgr(f.color.rgb)
                fl = cell.fill
                if fl is not None and fl.fill_type == "solid":
                    rgb = getattr(fl.start_color, "rgb", None)
                    cd.fill = hex_to_bgr(rgb) if isinstance(rgb, str) else None
                bd = cell.border
                for k in ("left", "right", "top", "bottom"):
                    side = getattr(bd, k)
                    if side is not None and side.style:
                        ls, wt = name_to_ls_weight(side.style)
                        col = hex_to_bgr(getattr(side.color, "rgb", None)) if side.color is not None else None
                        cd.borders[k] = (ls, col, wt)
                al = cell.alignment
                cd.halign = al.horizontal
                cd.wrap = al.wrap_text
                cd.locked = bool(cell.protection.locked)
                if (cd.value is not None or cd.formula is not None or cd.fill is not None
                        or cd.borders or cd.bold or cd.font_color is not None
                        or cd.number_format != "General" or cd.locked is not True):
                    sd.cells[(cell.row, cell.column)] = cd
        for letter, dim in ws.column_dimensions.items():
            try: idx = column_index_from_string(letter)
            except Exception: continue
            if dim.width: sd.col_width[idx] = dim.width
            if dim.hidden: sd.col_hidden.add(idx)
        for mc in ws.merged_cells.ranges:
            sd.merges.append((mc.min_row, mc.min_col, mc.max_row, mc.max_col))
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            ref = str(rng.sqref)
            for part in ref.split():
                m = re.match(r"\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", part)
                if not m: continue
                c1 = column_index_from_string(m.group(1)); r1 = int(m.group(2))
                c2 = column_index_from_string(m.group(3)) if m.group(3) else c1
                r2 = int(m.group(4)) if m.group(4) else r1
                for rule in rules:
                    if rule.type != "expression" or not rule.formula: continue
                    cr = CFRule(r1, c1, r2, c2, "=" + rule.formula[0])
                    dxf = rule.dxf
                    if dxf is not None:
                        if dxf.fill is not None:
                            rgb = getattr(dxf.fill, "bgColor", None)
                            rgb = getattr(rgb, "rgb", None)
                            if not isinstance(rgb, str):
                                rgb = getattr(getattr(dxf.fill, "fgColor", None), "rgb", None)
                            cr.fill = hex_to_bgr(rgb) if isinstance(rgb, str) else None
                        if dxf.border is not None:
                            for k in ("left", "right", "top", "bottom"):
                                side = getattr(dxf.border, k)
                                if side is not None and side.style:
                                    ls, wt = name_to_ls_weight(side.style)
                                    col = hex_to_bgr(getattr(side.color, "rgb", None)) if side.color is not None else None
                                    cr.borders[k] = (ls, col, wt)
                    sd.cf.append(cr)
        for dv in ws.data_validations.dataValidation:
            for part in str(dv.sqref).split():
                m = re.match(r"\$?([A-Z]+)\$?(\d+)$", part)
                if m:
                    sd.validations[(int(m.group(2)), column_index_from_string(m.group(1)))] = {
                        "type": dv.type, "op": dv.operator, "f1": dv.formula1,
                        "f2": dv.formula2, "title": dv.errorTitle, "msg": dv.error}
        if ws.freeze_panes:
            m = re.match(r"([A-Z]+)(\d+)", str(ws.freeze_panes))
            if m:
                sd.freeze = (int(m.group(2)) - 1, column_index_from_string(m.group(1)) - 1)
        sheets.append(sd)
    prot = getattr(wb, "security", None)
    return sheets, bool(prot and getattr(prot, "lockStructure", False))

# -------------------------------------------------------------- VBA objects

class VObj:
    type_name = "Object"
    def vba_get(self, name, args=(), kwargs=None):
        m = getattr(self, "g_" + name.lower(), None)
        if m is None:
            raise VBAError(438, "%s has no member %s" % (self.type_name, name))
        return m(list(args), dict(kwargs or {}))
    def vba_set(self, name, args, value):
        m = getattr(self, "s_" + name.lower(), None)
        if m is None:
            raise VBAError(438, "%s has no settable member %s" % (self.type_name, name))
        return m(list(args), value)

def _arg(args, i, default=MISSING):
    if i < len(args) and args[i] is not MISSING: return args[i]
    return default

def _kw(kwargs, name, default=MISSING):
    for k, v in kwargs.items():
        if k.lower() == name.lower(): return v
    return default

class RangeObj(VObj):
    type_name = "Range"
    def __init__(self, sheet, r1, c1, r2=None, c2=None):
        self.sh = sheet
        self.r1, self.c1 = r1, c1
        self.r2 = r2 if r2 is not None else r1
        self.c2 = c2 if c2 is not None else c1

    def cells_iter(self, create=True):
        if self.r2 >= MAXROW or self.c2 >= MAXCOL:
            for (r, c) in sorted(self.sh.cells):
                if self.r1 <= r <= self.r2 and self.c1 <= c <= self.c2:
                    yield r, c
            return
        for r in range(self.r1, self.r2 + 1):
            for c in range(self.c1, self.c2 + 1):
                yield r, c

    def _set_all(self, fn):
        if self.r2 >= MAXROW or self.c2 >= MAXCOL:
            for r, c in list(self.cells_iter()): fn(self.sh.cell(r, c))
            return True
        for r, c in self.cells_iter(): fn(self.sh.cell(r, c))
        return False

    # -- values
    def g_value(self, a, k):
        cd = self.sh.cells.get((self.r1, self.c1))
        if cd is None or (cd.value is None and cd.formula is None): return EMPTY
        if cd.value is None: return EMPTY
        return cd.value
    def s_value(self, a, v):
        if isinstance(v, str) and v.startswith("'"): v = v[1:]
        if v is EMPTY: v = None
        def setter(cd): cd.value = v; cd.formula = None
        self._set_all(setter)
    def g_formula(self, a, k):
        cd = self.sh.cells.get((self.r1, self.c1))
        if cd is None: return ""
        return cd.formula if cd.formula is not None else (vba_str(cd.value) if cd.value is not None else "")
    def s_formula(self, a, v):
        s = vba_str(v)
        def setter(cd):
            if s.startswith("="): cd.formula = s; cd.value = None
            else: cd.value = s; cd.formula = None
        self._set_all(setter)
    def g_text(self, a, k): return vba_str(self.g_value(a, k))
    def g_numberformat(self, a, k):
        cd = self.sh.cells.get((self.r1, self.c1))
        return cd.number_format if cd else "General"
    def s_numberformat(self, a, v):
        s = vba_str(v)
        self._set_all(lambda cd: setattr(cd, "number_format", s))

    # -- style
    def g_font(self, a, k): return FontObj(self)
    def g_interior(self, a, k): return InteriorObj(self)
    def g_borders(self, a, k):
        if a: return BorderObj(self, EDGE_KEY[int(num(a[0]))])
        return BordersCollection(self)
    def g_horizontalalignment(self, a, k):
        cd = self.sh.cells.get((self.r1, self.c1))
        return cd.halign if cd else None
    def s_horizontalalignment(self, a, v):
        nm = ALIGN.get(int(num(v)), None)
        self._set_all(lambda cd: setattr(cd, "halign", nm))
    def g_wraptext(self, a, k):
        cd = self.sh.cells.get((self.r1, self.c1))
        return bool(cd.wrap) if cd else False
    def s_wraptext(self, a, v):
        self._set_all(lambda cd: setattr(cd, "wrap", bool(v)))
    def g_locked(self, a, k): return self.sh.locked_at(self.r1, self.c1)
    def s_locked(self, a, v):
        b = bool(v)
        whole = (self.r2 >= MAXROW and self.c2 >= MAXCOL)
        if whole:
            self.sh.default_locked = b
            for cd in self.sh.cells.values(): cd.locked = None
            return
        self._set_all(lambda cd: setattr(cd, "locked", b))

    # -- geometry
    def g_count(self, a, k): return (self.r2 - self.r1 + 1) * (self.c2 - self.c1 + 1)
    def g_row(self, a, k): return self.r1
    def g_column(self, a, k): return self.c1
    def g_address(self, a, k):
        ra = bool(_arg(a, 0, True)); ca = bool(_arg(a, 1, True))
        def one(r, c):
            return "%s%s%s%d" % ("$" if ca else "", get_column_letter(c), "$" if ra else "", r)
        if (self.r1, self.c1) == (self.r2, self.c2): return one(self.r1, self.c1)
        return one(self.r1, self.c1) + ":" + one(self.r2, self.c2)
    def g_entirecolumn(self, a, k): return RangeObj(self.sh, 1, self.c1, MAXROW, self.c2)
    def g_entirerow(self, a, k): return RangeObj(self.sh, self.r1, 1, self.r2, MAXCOL)
    def g_columnwidth(self, a, k): return self.sh.col_width.get(self.c1, 8.43)
    def s_columnwidth(self, a, v):
        for c in range(self.c1, self.c2 + 1): self.sh.col_width[c] = float(num(v))
    def g_rowheight(self, a, k): return 13.5
    def s_rowheight(self, a, v): pass
    def g_hidden(self, a, k): return self.c1 in self.sh.col_hidden
    def s_hidden(self, a, v):
        for c in range(self.c1, self.c2 + 1):
            if bool(v): self.sh.col_hidden.add(c)
            else: self.sh.col_hidden.discard(c)
    def g_cells(self, a, k):
        if not a: return RangeObj(self.sh, self.r1, self.c1, self.r2, self.c2)
        r = self.r1 + int(num(a[0])) - 1; c = self.c1 + int(num(a[1])) - 1
        return RangeObj(self.sh, r, c)
    def g_rows(self, a, k):
        if not a: return self
        r = self.r1 + int(num(a[0])) - 1
        return RangeObj(self.sh, r, self.c1, r, self.c2)
    def g_columns(self, a, k):
        if not a: return self
        c = self.c1 + int(num(a[0])) - 1
        return RangeObj(self.sh, self.r1, c, self.r2, c)

    # -- actions
    def g_merge(self, a, k):
        self.sh.merges.append((self.r1, self.c1, self.r2, self.c2)); return EMPTY
    def g_select(self, a, k): return EMPTY
    def g_activate(self, a, k): return EMPTY
    def g_clearcontents(self, a, k):
        for r, c in list(self.cells_iter()):
            cd = self.sh.cells.get((r, c))
            if cd: cd.value = None; cd.formula = None
        return EMPTY
    def g_clear(self, a, k):
        whole = (self.r2 >= MAXROW and self.c2 >= MAXCOL)
        if whole:
            self.sh.cells.clear(); self.sh.merges = []
            self.sh.cf = []; self.sh.validations = {}
            self.sh.col_width = {}; self.sh.col_hidden = set()
        else:
            for r, c in list(self.cells_iter()):
                self.sh.cells.pop((r, c), None)
        return EMPTY
    def g_formatconditions(self, a, k): return FormatConditionsObj(self)
    def g_validation(self, a, k): return ValidationObj(self)
    def g_displayformat(self, a, k): return DisplayFormatObj(self.sh, self.r1, self.c1)

class FontObj(VObj):
    type_name = "Font"
    def __init__(self, rng): self.rng = rng
    def g_bold(self, a, k):
        cd = self.rng.sh.cells.get((self.rng.r1, self.rng.c1))
        return bool(cd.bold) if cd else False
    def s_bold(self, a, v): self.rng._set_all(lambda cd: setattr(cd, "bold", bool(v)))
    def g_color(self, a, k):
        cd = self.rng.sh.cells.get((self.rng.r1, self.rng.c1))
        return cd.font_color if cd and cd.font_color is not None else 0
    def s_color(self, a, v): self.rng._set_all(lambda cd: setattr(cd, "font_color", int(num(v))))
    def s_name(self, a, v): pass
    def s_size(self, a, v): pass

class InteriorObj(VObj):
    type_name = "Interior"
    def __init__(self, rng): self.rng = rng
    def g_color(self, a, k):
        cd = self.rng.sh.cells.get((self.rng.r1, self.rng.c1))
        return cd.fill if cd and cd.fill is not None else 16777215
    def s_color(self, a, v): self.rng._set_all(lambda cd: setattr(cd, "fill", int(num(v))))
    def g_colorindex(self, a, k):
        cd = self.rng.sh.cells.get((self.rng.r1, self.rng.c1))
        return -4142 if (cd is None or cd.fill is None) else 1
    def s_colorindex(self, a, v):
        if int(num(v)) == -4142: self.rng._set_all(lambda cd: setattr(cd, "fill", None))
    def g_pattern(self, a, k): return self.g_colorindex(a, k)
    def s_pattern(self, a, v): self.s_colorindex(a, v)

class BordersCollection(VObj):
    type_name = "Borders"
    def __init__(self, rng): self.rng = rng
    def g_item(self, a, k): return BorderObj(self.rng, EDGE_KEY[int(num(a[0]))])

class BorderObj(VObj):
    type_name = "Border"
    def __init__(self, rng, key): self.rng, self.key = rng, key
    def _cur(self):
        cd = self.rng.sh.cells.get((self.rng.r1, self.rng.c1))
        if cd is None: return (LS_NONE, None, W_THIN)
        return cd.borders.get(self.key, (LS_NONE, None, W_THIN))
    def _upd(self, idx, v):
        def setter(cd):
            cur = list(cd.borders.get(self.key, (LS_NONE, None, W_THIN)))
            cur[idx] = v
            cd.borders[self.key] = tuple(cur)
        self.rng._set_all(setter)
    def g_linestyle(self, a, k): return self._cur()[0]
    def s_linestyle(self, a, v): self._upd(0, int(num(v)))
    def g_color(self, a, k):
        c = self._cur()[1]
        return 0 if c is None else c
    def s_color(self, a, v): self._upd(1, int(num(v)))
    def g_weight(self, a, k): return self._cur()[2]
    def s_weight(self, a, v): self._upd(2, int(num(v)))

class DisplayFormatObj(VObj):
    type_name = "DisplayFormat"
    def __init__(self, sh, r, c):
        self.sh, self.r, self.c = sh, r, c
        self.fill, self.borders = effective_format(sh, r, c)
    def g_interior(self, a, k): return _DFInterior(self.fill)
    def g_borders(self, a, k): return _DFBorder(self.borders.get(EDGE_KEY[int(num(a[0]))]))

class _DFInterior(VObj):
    type_name = "Interior"
    def __init__(self, fill): self.fill = fill
    def g_color(self, a, k): return 16777215 if self.fill is None else self.fill
    def g_colorindex(self, a, k): return -4142 if self.fill is None else 1

class _DFBorder(VObj):
    type_name = "Border"
    def __init__(self, tup): self.tup = tup or (LS_NONE, None, W_THIN)
    def g_linestyle(self, a, k): return self.tup[0]
    def g_color(self, a, k): return 0 if self.tup[1] is None else self.tup[1]
    def g_weight(self, a, k): return self.tup[2]

class FormatConditionsObj(VObj):
    type_name = "FormatConditions"
    def __init__(self, rng): self.rng = rng
    def g_add(self, a, k):
        formula = None
        for x in a[2:]:
            if isinstance(x, str): formula = x; break
        if formula is None: formula = _kw(k, "Formula1", "")
        rule = CFRule(self.rng.r1, self.rng.c1, self.rng.r2, self.rng.c2, vba_str(formula))
        self.rng.sh.cf.append(rule)
        return FormatConditionObj(rule)
    def g_delete(self, a, k):
        sh = self.rng.sh
        if self.rng.r2 >= MAXROW and self.rng.c2 >= MAXCOL:
            sh.cf = []
        else:
            sh.cf = [x for x in sh.cf if not (x.r1 >= self.rng.r1 and x.r2 <= self.rng.r2
                                              and x.c1 >= self.rng.c1 and x.c2 <= self.rng.c2)]
        return EMPTY
    def g_count(self, a, k): return len(self.rng.sh.cf)

class FormatConditionObj(VObj):
    type_name = "FormatCondition"
    def __init__(self, rule): self.rule = rule
    def g_interior(self, a, k): return _CFInterior(self.rule)
    def g_borders(self, a, k): return _CFBorder(self.rule, EDGE_KEY[int(num(a[0]))])
    def g_font(self, a, k): return _CFFont(self.rule)
    def s_stopiftrue(self, a, v): pass

class _CFInterior(VObj):
    type_name = "Interior"
    def __init__(self, rule): self.rule = rule
    def s_color(self, a, v): self.rule.fill = int(num(v))
    def g_color(self, a, k): return self.rule.fill if self.rule.fill is not None else 16777215

class _CFFont(VObj):
    type_name = "Font"
    def __init__(self, rule): self.rule = rule
    def s_bold(self, a, v): pass
    def s_color(self, a, v): pass

class _CFBorder(VObj):
    type_name = "Border"
    def __init__(self, rule, key): self.rule, self.key = rule, key
    def _upd(self, idx, v):
        cur = list(self.rule.borders.get(self.key, (LS_NONE, None, W_THIN)))
        cur[idx] = v
        self.rule.borders[self.key] = tuple(cur)
    def s_linestyle(self, a, v): self._upd(0, int(num(v)))
    def s_color(self, a, v): self._upd(1, int(num(v)))
    def s_weight(self, a, v): self._upd(2, int(num(v)))
    def g_linestyle(self, a, k): return self.rule.borders.get(self.key, (LS_NONE, None, W_THIN))[0]

class ValidationObj(VObj):
    type_name = "Validation"
    def __init__(self, rng): self.rng = rng
    def _key(self): return (self.rng.r1, self.rng.c1)
    def g_add(self, a, k):
        t = int(num(_arg(a, 0, 4)))
        op = int(num(_arg(a, 2, 1)))
        f1, f2 = _arg(a, 3, ""), _arg(a, 4, "")
        def fx(v):
            if v is MISSING or v is EMPTY: return None
            if isinstance(v, VDate):
                d = EPOCH + datetime.timedelta(days=int(v))
                return d.strftime("%Y-%m-%d")
            return vba_str(v)
        self.rng.sh.validations[self._key()] = {
            "type": {4: "date", 1: "whole", 7: "custom"}.get(t, "date"),
            "op": {1: "between"}.get(op, "between"),
            "f1": fx(f1), "f2": fx(f2), "title": None, "msg": None}
        return EMPTY
    def g_delete(self, a, k):
        if self.rng.r2 >= MAXROW and self.rng.c2 >= MAXCOL:
            self.rng.sh.validations = {}
        else:
            for r, c in list(self.rng.cells_iter()):
                self.rng.sh.validations.pop((r, c), None)
        return EMPTY
    def _set(self, field, v):
        d = self.rng.sh.validations.get(self._key())
        if d is not None: d[field] = vba_str(v)
    def s_errortitle(self, a, v): self._set("title", v)
    def s_errormessage(self, a, v): self._set("msg", v)
    def s_showerror(self, a, v): pass
    def s_ignoreblank(self, a, v): pass
    def g_type(self, a, k):
        d = self.rng.sh.validations.get(self._key())
        if d is None: raise VBAError(1004, "no validation")
        return {"date": 4, "whole": 1, "custom": 7}.get(d["type"], 4)

# -------------------------------------------------------------- sheets / books / app

A1_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?$")

class WorksheetObj(VObj):
    type_name = "Worksheet"
    def __init__(self, app, wb, sd): self.app, self.wb, self.sd = app, wb, sd
    def g_name(self, a, k): return self.sd.name
    def s_name(self, a, v): self.sd.name = vba_str(v)
    def g_parent(self, a, k): return self.wb
    def g_cells(self, a, k):
        if not a: return RangeObj(self.sd, 1, 1, MAXROW, MAXCOL)
        return RangeObj(self.sd, int(num(a[0])), int(num(a[1])))
    def g_range(self, a, k):
        if len(a) >= 2 and isinstance(a[0], RangeObj):
            x, y = a[0], a[1]
            return RangeObj(self.sd, min(x.r1, y.r1), min(x.c1, y.c1),
                            max(x.r2, y.r2), max(x.c2, y.c2))
        if isinstance(a[0], RangeObj): return a[0]
        m = A1_RE.match(vba_str(a[0]).upper())
        if not m: raise VBAError(1004, "bad address %r" % (a[0],))
        c1 = column_index_from_string(m.group(1)); r1 = int(m.group(2))
        c2 = column_index_from_string(m.group(3)) if m.group(3) else c1
        r2 = int(m.group(4)) if m.group(4) else r1
        return RangeObj(self.sd, r1, c1, r2, c2)
    def g_columns(self, a, k):
        if not a: return RangeObj(self.sd, 1, 1, MAXROW, MAXCOL)
        v = a[0]
        c = column_index_from_string(v.upper()) if isinstance(v, str) else int(num(v))
        return RangeObj(self.sd, 1, c, MAXROW, c)
    def g_rows(self, a, k):
        if not a: return RangeObj(self.sd, 1, 1, MAXROW, MAXCOL)
        r = int(num(a[0]))
        return RangeObj(self.sd, r, 1, r, MAXCOL)
    def g_usedrange(self, a, k):
        mr, mc = self.sd.max_used()
        return RangeObj(self.sd, 1, 1, max(1, mr), max(1, mc))
    def g_visible(self, a, k):
        return {"visible": -1, "hidden": 0, "veryHidden": 2}[self.sd.visible]
    def s_visible(self, a, v):
        self.sd.visible = {-1: "visible", 0: "hidden", 2: "veryHidden"}[int(num(v))]
    def g_activate(self, a, k):
        self.app.active_sheet = self; self.app.active_wb = self.wb; return EMPTY
    def g_select(self, a, k): return self.g_activate(a, k)
    def g_protect(self, a, k): self.sd.protected = True; return EMPTY
    def g_unprotect(self, a, k): self.sd.protected = False; return EMPTY
    def g_protectcontents(self, a, k): return self.sd.protected
    def g_shapes(self, a, k): return ShapesObj(self.sd, a)
    def g_copy(self, a, k):
        if not a and not k:
            wb = WorkbookObj(self.app, [self.sd.clone()], name="Book%d" % (len(self.app.workbooks) + 2))
            self.app.workbooks.append(wb)
            self.app.active_wb = wb
            self.app.active_sheet = wb.sheet_objs()[0]
            return EMPTY
        raise VBAError(1004, "Worksheet.Copy with arguments is not supported")

class WorksheetsObj(VObj):
    type_name = "Sheets"
    def __init__(self, wb): self.wb = wb
    def g_item(self, a, k): return self.wb.find_sheet(a[0])
    def g_count(self, a, k): return len(self.wb.sheets)
    def g_add(self, a, k):
        before = _arg(a, 0, _kw(k, "Before"))
        after = _arg(a, 1, _kw(k, "After"))
        sd = SheetData("Sheet%d" % (len(self.wb.sheets) + 1))
        idx = len(self.wb.sheets)
        if isinstance(after, WorksheetObj): idx = self.wb.sheets.index(after.sd) + 1
        elif isinstance(before, WorksheetObj): idx = self.wb.sheets.index(before.sd)
        self.wb.sheets.insert(idx, sd)
        return WorksheetObj(self.wb.app, self.wb, sd)

class WorkbookObj(VObj):
    type_name = "Workbook"
    def __init__(self, app, sheets, name="Book1", path=None, protect=False):
        self.app, self.sheets, self.name, self.path = app, sheets, name, path
        self.protect_structure = protect
        self.saved = True
    def sheet_objs(self): return [WorksheetObj(self.app, self, sd) for sd in self.sheets]
    def find_sheet(self, key):
        if isinstance(key, str):
            for sd in self.sheets:
                if sd.name == key: return WorksheetObj(self.app, self, sd)
            raise VBAError(9, "Subscript out of range (sheet %r)" % key)
        i = int(num(key))
        if 1 <= i <= len(self.sheets):
            return WorksheetObj(self.app, self, self.sheets[i - 1])
        raise VBAError(9, "Subscript out of range (sheet %d)" % i)
    def g_worksheets(self, a, k):
        return self.find_sheet(a[0]) if a else WorksheetsObj(self)
    g_sheets = g_worksheets
    def g_name(self, a, k):
        return self.name if not self.path else self.path.replace("\\", "/").split("/")[-1]
    def g_fullname(self, a, k): return self.path or self.name
    def g_saved(self, a, k): return self.saved
    def s_saved(self, a, v): self.saved = bool(v)
    def g_protect(self, a, k):
        st = _arg(a, 1, _kw(k, "Structure", True))
        self.protect_structure = bool(st) if st is not MISSING else True
        return EMPTY
    def g_unprotect(self, a, k): self.protect_structure = False; return EMPTY
    def g_saveas(self, a, k):
        path = _arg(a, 0, _kw(k, "Filename"))
        save_workbook(self, vba_str(path))
        self.path = vba_str(path); self.saved = True
        return EMPTY
    def g_save(self, a, k):
        if self.path: save_workbook(self, self.path)
        return EMPTY
    def g_close(self, a, k):
        if self in self.app.workbooks: self.app.workbooks.remove(self)
        if self.app.active_wb is self:
            self.app.active_wb = self.app.workbooks[-1] if self.app.workbooks else NOTHING
        return EMPTY
    def g_activate(self, a, k): self.app.active_wb = self; return EMPTY

class WorkbooksObj(VObj):
    type_name = "Workbooks"
    def __init__(self, app): self.app = app
    def g_item(self, a, k):
        i = int(num(a[0])); return self.app.workbooks[i - 1]
    def g_count(self, a, k): return len(self.app.workbooks)
    def g_add(self, a, k):
        wb = WorkbookObj(self.app, [SheetData("Sheet1")],
                         name="Book%d" % (len(self.app.workbooks) + 2))
        self.app.workbooks.append(wb); self.app.active_wb = wb
        return wb
    def g_open(self, a, k):
        path = vba_str(_arg(a, 0, _kw(k, "Filename")))
        sheets, prot = load_workbook_data(path)
        wb = WorkbookObj(self.app, sheets, path=path, protect=prot)
        self.app.workbooks.append(wb); self.app.active_wb = wb
        return wb

class WindowObj(VObj):
    type_name = "Window"
    def __init__(self, app): self.app = app; self.split_row = 0; self.split_col = 0
    def s_splitrow(self, a, v): self.split_row = int(num(v))
    def g_splitrow(self, a, k): return self.split_row
    def s_splitcolumn(self, a, v): self.split_col = int(num(v))
    def g_splitcolumn(self, a, k): return self.split_col
    def s_freezepanes(self, a, v):
        sh = self.app.active_sheet
        if sh is None or sh is NOTHING: return
        sh.sd.freeze = (self.split_row, self.split_col) if bool(v) else None
    def g_freezepanes(self, a, k):
        sh = self.app.active_sheet
        return bool(sh and sh.sd.freeze)

class ShapesObj(VObj):
    type_name = "Shapes"
    def __init__(self, sd, a): self.sd = sd; self.a = a
    def g_count(self, a, k): return len(self.sd.shapes)
    def g_item(self, a, k): return self.sd.shapes[int(num(a[0])) - 1]
    def g_addformcontrol(self, a, k):
        s = ShapeObj(); self.sd.shapes.append(s); return s

class ShapeObj(VObj):
    type_name = "Shape"
    def __init__(self): self.name = ""; self.onaction = ""; self.text = ""
    def g_name(self, a, k): return self.name
    def s_name(self, a, v): self.name = vba_str(v)
    def s_onaction(self, a, v): self.onaction = vba_str(v)
    def g_textframe(self, a, k): return _TextFrame(self)

class _TextFrame(VObj):
    type_name = "TextFrame"
    def __init__(self, shp): self.shp = shp
    def g_characters(self, a, k): return _Characters(self.shp)

class _Characters(VObj):
    type_name = "Characters"
    def __init__(self, shp): self.shp = shp
    def s_text(self, a, v): self.shp.text = vba_str(v)
    def g_text(self, a, k): return self.shp.text

class ApplicationObj(VObj):
    type_name = "Application"
    def __init__(self):
        self.workbooks = []
        self.active_wb = NOTHING
        self.active_sheet = None
        self.window = WindowObj(self)
        self.screen_updating = True
        self.display_alerts = True
        self.calculation = -4105
        self.picked_file = None
    def g_workbooks(self, a, k):
        wbs = WorkbooksObj(self)
        return wbs.g_item(a, k) if a else wbs
    def g_activeworkbook(self, a, k): return self.active_wb
    def g_thisworkbook(self, a, k): return self.active_wb
    def g_activesheet(self, a, k): return self.active_sheet or NOTHING
    def g_activewindow(self, a, k): return self.window
    def g_screenupdating(self, a, k): return self.screen_updating
    def s_screenupdating(self, a, v): self.screen_updating = bool(v)
    def g_displayalerts(self, a, k): return self.display_alerts
    def s_displayalerts(self, a, v): self.display_alerts = bool(v)
    def g_calculation(self, a, k): return self.calculation
    def s_calculation(self, a, v): self.calculation = int(num(v))
    def g_getopenfilename(self, a, k):
        return self.picked_file if self.picked_file else False
    def g_union(self, a, k):
        rs = [x for x in a if isinstance(x, RangeObj)]
        sh = rs[0].sh
        return RangeObj(sh, min(r.r1 for r in rs), min(r.c1 for r in rs),
                        max(r.r2 for r in rs), max(r.c2 for r in rs))
    def g_worksheetfunction(self, a, k): return _WsFunc()

class _WsFunc(VObj):
    type_name = "WorksheetFunction"
    def g_max(self, a, k): return max(num(x) for x in a)
    def g_min(self, a, k): return min(num(x) for x in a)
