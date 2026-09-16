# -*- coding: utf-8 -*-
"""設計A 8章の受け入れ試験 1〜6。

納品する .bas をそのまま実行する。合否はログと出力 xlsx の中身で判定する。
"""
import os, sys, shutil, csv, io, datetime, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import harness, vbaint, excel_mock
import openpyxl
from openpyxl.utils import get_column_letter

SAMPLE = sys.argv[1]
OUTDIR = sys.argv[2]
RESULTS = []

def passed(log):
    lines = [l.strip() for l in log.splitlines()]
    return ("合格" in lines) and ("不合格" not in lines)

def record(no, title, ok, detail):
    RESULTS.append((no, title, ok, detail))
    print("%s 試験%s %s" % ("PASS" if ok else "FAIL", no, title))
    for line in detail.splitlines():
        print("        " + line)

def fresh(tag):
    d = os.path.join(OUTDIR, tag)
    if os.path.exists(d): shutil.rmtree(d)
    os.makedirs(d)
    shutil.copy(SAMPLE, d)
    return d, os.path.join(d, os.path.basename(SAMPLE))

def run(csvpath, d1, d2):
    I = harness.build()
    mod, proc = I.find_proc("Convert")
    msg = vbaint.vba_str(I.call_proc(mod, proc, [csvpath, d1, d2]))
    log = open(os.path.join(os.path.dirname(csvpath), "変換ログ.txt"), encoding="utf-8-sig").read()
    out = msg.split("→ ", 1)[1].strip() if "→ " in msg else None
    return msg, log, out

D = harness.date

# ---------------------------------------------------------------- 1
d, csvp = fresh("t1")
msg, log, out = run(csvp, D(2026, 9, 1), D(2026, 10, 10))
ngs = [l for l in log.splitlines() if l.startswith("NG")]
oks = [l for l in log.splitlines() if l.startswith("OK")]
record("1", "Start=2026/09/01 End=2026/10/10 で 6章の検査が全件合格",
       passed(log) and not ngs and len(oks) >= 12,
       "検査 OK %d 件 / NG %d 件\n出力: %s" % (len(oks), len(ngs), os.path.basename(out or "")))
T1_OUT = out

# ---------------------------------------------------------------- 2
d, csvp = fresh("t2")
msg, log, out = run(csvp, D(2026, 9, 1), D(2026, 9, 30))
line = [l for l in log.splitlines() if l.startswith("描画:")][0]
n_out = int(line.split("期間外で除外:")[1].split("件")[0].strip())
oct_only = []
with open(SAMPLE, encoding="utf-8-sig") as f:
    rows = list(csv.reader(f))
H = {h: i for i, h in enumerate(rows[2])}
for r in rows[3:]:
    if r[H["開始日"]][:10] > "2026-09-30":
        import json
        oct_only.append(json.loads(r[H["工程線名"]])[0]["name"])
record("2", "Start=2026/09/01 End=2026/09/30 で期間外除外がログに出る",
       n_out == len(oct_only) and n_out > 0 and "不合格" not in log,
       "%s\n10月開始の工程: %s (%d 件) → 除外数と一致: %s"
       % (line, ", ".join(oct_only), len(oct_only), n_out == len(oct_only)))

# ---------------------------------------------------------------- 3
d, csvp = fresh("t3")
msg, log, out = run(csvp, D(2026, 9, 15), D(2026, 10, 10))
wb = openpyxl.load_workbook(out); ws = wb["工程表"]
sheets, _ = excel_mock.load_workbook_data(out)
sd = [s for s in sheets if s.name == "工程表"][0]
clipped = []
for r in rows[3:]:
    s0, e0 = r[H["開始日"]][:10], r[H["終了日"]][:10]
    if s0 < "2026-09-15" <= e0:
        import json
        clipped.append((json.loads(r[H["工程線名"]])[0]["name"], r[H["工程ID"]]))
# 出力上で、クリップされた工程のバーの最初の列が 9/15 か
idcol = 6
while ws.cell(2, idcol).value is not None: idcol += 1
rowof = {}
for rr in range(4, ws.max_row + 1):
    v = ws.cell(rr, idcol).value
    if v: rowof[str(v)] = rr
bad = []
for name, pid in clipped:
    rr = rowof.get(pid)
    first = None
    for c in range(6, 6 + 26):
        fill, bd = excel_mock.effective_format(sd, rr, c)
        painted = fill is not None or (bd.get("bottom") and bd["bottom"][0] != -4142)
        if painted:
            first = c; break
    dt = ws.cell(2, first).value if first else None
    if not (dt and dt.strftime("%Y/%m/%d") == "2026/09/15"):
        bad.append("%s -> %s" % (name, dt))
record("3", "Start=2026/09/15 でクリップされる工程の開始列が 9/15",
       len(clipped) > 0 and not bad and "不合格" not in log,
       "クリップ対象 %d 件: %s\n開始列が 9/15 でないもの: %s"
       % (len(clipped), ", ".join(n for n, _ in clipped), bad or "なし"))

# ---------------------------------------------------------------- 4
wb4 = openpyxl.load_workbook(T1_OUT, keep_vba=False)
ws4 = wb4["工程表"]
import zipfile
z = zipfile.ZipFile(T1_OUT)
has_vba = any("vbaProject" in n for n in z.namelist())
unlocked, locked_ok = [], True
idcol4 = 6
while ws4.cell(2, idcol4).value is not None: idcol4 += 1
proc_rows = [r for r in range(4, ws4.max_row + 1) if ws4.cell(r, idcol4).value]
for r in proc_rows:
    if ws4.cell(r, 3).protection.locked or ws4.cell(r, 4).protection.locked:
        locked_ok = False
    for c in [1, 2, 5, 6, idcol4]:
        if not ws4.cell(r, c).protection.locked: unlocked.append((r, c))
record("4", "マクロ無効の Excel で C/D だけ編集できる",
       (not has_vba) and ws4.protection.sheet and locked_ok and not unlocked
       and bool(wb4.security and wb4.security.lockStructure),
       "マクロ痕跡 (vbaProject): %s / シート保護: %s / ブック構造保護: %s\n"
       "工程行 %d 件すべてで C・D が編集可、A・B・E・日付列・ID列はロック: %s"
       % (has_vba, ws4.protection.sheet,
          bool(wb4.security and wb4.security.lockStructure), len(proc_rows),
          locked_ok and not unlocked))

# ---------------------------------------------------------------- 6
d, csvp = fresh("t6")
with open(SAMPLE, "rb") as f: raw = f.read().decode("utf-8-sig")
rr = list(csv.reader(io.StringIO(raw, newline="")))
rr.append(list(rr[3]))                      # 工程行を 1 本そのまま複製 -> 24 本
buf = io.StringIO()
w = csv.writer(buf, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
for row in rr: w.writerow(row)
with open(csvp, "wb") as f: f.write(b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8"))
msg, log, out = run(csvp, D(2026, 9, 1), D(2026, 10, 10))
loaded = [l for l in log.splitlines() if l.startswith("読み込み工程数")]
drawn = [l for l in log.splitlines() if l.startswith("描画:")]
n_drawn = int(drawn[0].split("描画:")[1].split("件")[0].strip()) if drawn else -1
ngs6 = [l for l in log.splitlines() if l.startswith("NG")]
record("6", "工程行を複製した 24 本の CSV でも停止せず 24 本描く",
       "24 件" in (loaded[0] if loaded else "") and n_drawn == 24 and not ngs6
       and passed(log),
       "%s\n%s\nNG: %d 件" % (loaded[0] if loaded else "(なし)", drawn[0] if drawn else "(なし)", len(ngs6)))

print()
print("=" * 64)
for no, title, ok, _ in RESULTS:
    print("  試験%s %-46s %s" % (no, title[:46], "合格" if ok else "不合格"))
print("  試験5 スマホ表示の確認 %-34s %s" % ("", "別スクリプト (phone_render.py)"))
print("=" * 64)
sys.exit(0 if all(r[2] for r in RESULTS) else 1)
