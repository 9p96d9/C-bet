# -*- coding: utf-8 -*-
"""プロジェクトG変換ツール.xlsm を組み立てる。

openpyxl で 4 シートのブックを作り、zip を開いて vbaProject.bin と
関連する content-types / relationship を差し込む (openpyxl 単体では
マクロ有効ブックを新規に作れないため)。
"""
import os, sys, shutil, zipfile, re, datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_vba

TOOL_VERSION = "A-1.0.0"

SHEETS = [("00_Control", "Sheet1"), ("T09_Audit", "Sheet2"),
          ("T10_Layout", "Sheet3"), ("使い方", "Sheet4")]

HOWTO = [
    ("プロジェクトG 工程表変換ツール（設計A 往路）使い方", True),
    ("", False),
    ("■ 配布物を開くまで", True),
    ("1. 受け取った ZIP を右クリック → プロパティ → 全般タブ下部の「セキュリティ:」で", False),
    ("   「許可する」にチェック → OK。（Mark of the Web の解除。これをしないと", False),
    ("   展開した xlsm が保護ビューで開き、マクロが動きません）", False),
    ("2. ZIP を展開する。展開後のフォルダーから xlsm を開く。", False),
    ("3. 「コンテンツの有効化」を押してマクロを有効にする。", False),
    ("", False),
    ("■ 変換の手順", True),
    ("4. 00_Control シートの「CSV 選択」ボタンで プロジェクトG の CSV を選ぶ。", False),
    ("   （ボタンが出ていないときは Alt+F8 → SelectCsv → 実行）", False),
    ("5. Start / End を確認する。既定は CSV の「工程表の期間」の先頭から 40 日。", False),
    ("   Start ≤ End、かつ「工程表の期間」の中に収まっている必要がある。", False),
    ("6. 「変換実行」ボタンを押す。（Alt+F8 → RunConversion でも同じ）", False),
    ("7. CSV と同じフォルダーに次の 2 つができる。", False),
    ("     ・<CSV名>_<Start>-<End>.xlsx … 業者に送るファイル（マクロなし）", False),
    ("     ・変換ログ.txt … 読み込み件数・除外件数・機械検査の結果", False),
    ("8. 変換ログ.txt の末尾が「合格」であることを確認してから業者に送る。", False),
    ("   「不合格」のときは NG 行に理由が出ている。xlsx は消さずに残してある。", False),
    ("", False),
    ("■ 出力 xlsx について（業者向けの説明）", True),
    ("・直してよいのは C 列（開始日）と D 列（終了日）だけ。他はロックしてある。", False),
    ("・日付を直すとバーが自動で動く（条件付き書式で描いているため）。", False),
    ("・行の挿入・削除・並べ替えはできない。プロジェクトG の行番号と対応が取れなくなるため。", False),
    ("・スマホ / タブレットの Excel でも開ける。マクロは入っていない。", False),
    ("", False),
    ("■ 仕様上の注意", True),
    ("・S / M / L は プロジェクトG 上のバーの高さの違い。Excel では同じ高さで描く。", False),
    ("・折れ線・斜行・crank / gate の縦線・ノードの丸・関係線は描かない。", False),
    ("  xlsx は業者が日程を直すためのガント投影であり、意図的に落としている。", False),
    ("・終了行が開始行と違う工程は、終了行側に ▼ を置いて行き先だけ示す。", False),
    ("・土日は灰色。祝日は v1 では扱わない（プロジェクトG の画面では祝日も灰色になる）。", False),
    ("・0.5 日・詳細工程列・復路は v1 では扱わない。", False),
    ("・同じ行に 2 本以上の工程が始まる場合、2 本目以降は次の空行に落とし、", False),
    ("  変換ログに「警告 行衝突」として出す。上書きはしない。", False),
]

def build_workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, code in SHEETS:
        ws = wb.create_sheet(name)
        ws.sheet_properties.codeName = code
    wb.code_name = "ThisWorkbook"

    ctl = wb["00_Control"]
    ctl["A1"] = "プロジェクトG 工程表変換ツール（設計A 往路）"
    ctl["A1"].font = Font(bold=True, size=14)
    ctl["A3"] = "CSV パス"
    ctl["A5"] = "Start"
    ctl["D5"] = "End"
    ctl["A7"] = "※ Start / End は「工程表の期間」の中で指定する。既定は先頭から 40 日。"
    ctl["A9"] = "結果"
    ctl["A11"] = "ボタンが出ていないときは Alt+F8 から SelectCsv / RunConversion を実行する。"
    ctl["A12"] = "詳しい手順は「使い方」シート。"
    ctl["A14"] = "版"
    ctl["B14"] = TOOL_VERSION
    for a in ("A3", "A5", "D5", "A9", "A14"):
        ctl[a].font = Font(bold=True)
    fill = PatternFill("solid", start_color="FFF6E0", end_color="FFF6E0")
    thin = Side(style="thin", color="BFBFBF")
    for a in ("B3", "B5", "E5"):
        ctl[a].fill = fill
        ctl[a].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctl["B5"].number_format = "yyyy/mm/dd"
    ctl["E5"].number_format = "yyyy/mm/dd"
    ctl["B9"].alignment = Alignment(wrap_text=False)
    ctl.column_dimensions["A"].width = 12
    ctl.column_dimensions["B"].width = 58
    ctl.column_dimensions["C"].width = 3
    ctl.column_dimensions["D"].width = 8
    ctl.column_dimensions["E"].width = 16

    hlp = wb["使い方"]
    for i, (text, bold) in enumerate(HOWTO, start=1):
        hlp.cell(i, 1).value = text
        if bold:
            hlp.cell(i, 1).font = Font(bold=True)
    hlp.column_dimensions["A"].width = 100

    wb["T09_Audit"]["A1"] = "（変換実行で自動生成されます）"
    wb["T10_Layout"]["A1"] = "（変換実行で自動生成されます）"
    wb.active = 0
    wb.save(path)

VBA_REL = ('<Relationship Id="rIdVbaProject" '
           'Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
           'Target="vbaProject.bin"/>')

def inject_vba(xlsx_path, xlsm_path, vba_bin):
    zin = zipfile.ZipFile(xlsx_path)
    items = {n: zin.read(n) for n in zin.namelist()}
    zin.close()

    ct = items["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    if 'Extension="bin"' not in ct:
        ct = ct.replace("<Types ", "<Types ", 1)
        ct = re.sub(r"(<Types[^>]*>)",
                    r'\1<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>',
                    ct, count=1)
    items["[Content_Types].xml"] = ct.encode("utf-8")

    rels = items["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rels = rels.replace("</Relationships>", VBA_REL + "</Relationships>")
    items["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    items["xl/vbaProject.bin"] = vba_bin

    with zipfile.ZipFile(xlsm_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in ["[Content_Types].xml"] + [n for n in items if n != "[Content_Types].xml"]:
            zout.writestr(name, items[name])

def main():
    srcdir, outdir = sys.argv[1], sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    tmp_xlsx = os.path.join(outdir, "_tool_tmp.xlsx")
    out = os.path.join(outdir, "プロジェクトG変換ツール.xlsm")
    build_workbook(tmp_xlsx)
    vba, mods = build_vba.build(srcdir, build_vba.DEFAULT_DOCS)
    inject_vba(tmp_xlsx, out, vba)
    os.remove(tmp_xlsx)
    print("wrote %s (%d bytes)" % (out, os.path.getsize(out)))
    print("modules: %s" % ", ".join(m[0] for m in mods))

main()
