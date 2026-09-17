#!/usr/bin/env python3
"""
出力 xlsx を ExcelJS とは別の実装（openpyxl）で読み直して検査する。
ExcelJS が自分の書いたものを読み返すだけでは「相互に間違っている」を
見逃すため、独立した読み手を 1 つ挟む。期待値は元 CSV から直接作る。
"""
import csv, io, json, re, sys, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "fixtures" / "代替サンプル工程表.csv"
XLSX = sys.argv[1] if len(sys.argv) > 1 else str(
    ROOT / "out" / "代替サンプル工程表_20260901-20261010.xlsx")
START = datetime.date(2026, 9, 1)
END = datetime.date(2026, 10, 10)

fails = []
lines = []


def chk(ok, label, detail=""):
    if not ok:
        fails.append(label)
    lines.append(f"{'OK' if ok else 'NG'}\t{label}\t{detail}")


def as_date(v):
    """openpyxl は numFmt 'aaa' を日付書式と見なさないのでシリアル値のまま返す。
    Excel 自身は 'aaa' を曜日として描画するので、格納値は正しい。
    ここではどちらの形でも日付として読む（1900 日付システム）。"""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    return None


# ---- 元 CSV を素直に読む（ツールとは別経路） ----
raw = CSV_PATH.read_text(encoding="utf-8-sig")
rows = list(csv.reader(io.StringIO(raw)))
headers = rows[2]
H = {name: i for i, name in enumerate(headers)}
procs = []
for r in rows[3:]:
    if not any(c.strip() for c in r):
        continue
    name = json.loads(r[H["工程線名"]])[0]["name"]
    s = datetime.date.fromisoformat(r[H["開始日"]][:10])
    e = datetime.date.fromisoformat(r[H["終了日"]][:10])
    procs.append({
        "id": r[H["工程ID"]], "name": name, "start": s, "end": e,
        "color": r[H["工程線の色"]], "raw": r,
    })
procs = [p for p in procs if p["end"] >= START and p["start"] <= END]

wb = load_workbook(XLSX)
chk(wb.sheetnames == ["T10_Layout", "_data", "使い方"], "シート名と並び", str(wb.sheetnames))
ws = wb["T10_Layout"]
wd = wb["_data"]
chk(wd.sheet_state == "hidden", "_data が非表示", wd.sheet_state)
chk(ws.protection.sheet is True, "T10_Layout がシート保護されている", str(ws.protection.sheet))

days = (END - START).days + 1
COL0 = 6  # F 列
last = COL0 + days - 1

# 行 2 の日付、行 3 の曜日、土日の薄灰
bad_day = bad_wk = bad_fill = 0
for n in range(days):
    d = START + datetime.timedelta(days=n)
    c2 = ws.cell(row=2, column=COL0 + n)
    c3 = ws.cell(row=3, column=COL0 + n)
    v2 = as_date(c2.value)
    v3 = as_date(c3.value)
    if v2 != d:
        bad_day += 1
    if v3 != d:
        bad_wk += 1
    shaded = (c2.fill is not None and c2.fill.fgColor is not None
              and str(c2.fill.fgColor.rgb or "").upper().endswith("E8E8E8"))
    if (d.weekday() >= 5) != shaded:
        bad_fill += 1
chk(bad_day == 0, "行 2 の日付が 1 日ずつ一致", f"{days} 列 / NG {bad_day}")
chk(bad_wk == 0, "行 3 が同じ日付を指す", f"NG {bad_wk}")
chk(bad_fill == 0, "土日列が薄灰", f"NG {bad_fill}")
chk(ws.cell(row=2, column=COL0).number_format == "d", "行 2 の書式 d",
    ws.cell(row=2, column=COL0).number_format)
chk(ws.cell(row=3, column=COL0).number_format == "aaa", "行 3 の書式 aaa",
    ws.cell(row=3, column=COL0).number_format)

# 月見出しの結合
merges = {str(m) for m in ws.merged_cells.ranges}
months = sorted({(START + datetime.timedelta(days=n)).month for n in range(days)})
chk(len(merges) == len(months), "月見出しの結合が月数だけある", f"{sorted(merges)}")

# 1 行 1 工程・C/D・保護
bad_name = bad_date = bad_lock = 0
for i, p in enumerate(procs):
    r = 5 + i
    if ws.cell(row=r, column=2).value != p["name"]:
        bad_name += 1
    cs, ce = ws.cell(row=r, column=3), ws.cell(row=r, column=4)
    vs, ve = as_date(cs.value), as_date(ce.value)
    if vs != p["start"] or ve != p["end"]:
        bad_date += 1
    if cs.protection.locked is not False or ce.protection.locked is not False:
        bad_lock += 1
    if ws.cell(row=r, column=1).protection.locked is False:
        bad_lock += 1
chk(bad_name == 0, "1 行 1 工程で工程線名が一致", f"{len(procs)} 行 / NG {bad_name}")
chk(bad_date == 0, "C/D が元 CSV の開始日・終了日と一致", f"NG {bad_date}")
chk(bad_lock == 0, "C/D だけ編集可、A 列は保護", f"NG {bad_lock}")

# 条件付き書式
cfs = list(ws.conditional_formatting)
chk(len(cfs) == len(procs), "条件付き書式が工程数だけある", f"{len(cfs)} / {len(procs)}")
bad_cf = []
by_ref = {str(cf.sqref): cf for cf in cfs}
for i, p in enumerate(procs):
    r = 5 + i
    ref = f"F{r}:{ws.cell(row=r, column=last).column_letter}{r}"
    cf = by_ref.get(ref)
    if cf is None:
        bad_cf.append(f"{p['id']}:ref無し({ref})")
        continue
    rule = cf.rules[0]
    want = f"AND(F$2>=$C{r},F$2<=$D{r})"
    got = (rule.formula[0] if rule.formula else "").lstrip("=")
    rgb = str(rule.dxf.fill.fgColor.rgb or "").upper() if rule.dxf and rule.dxf.fill else ""
    want_rgb = "FF" + p["color"].lstrip("#").upper()
    if rule.type != "expression" or got != want or rgb != want_rgb:
        bad_cf.append(f"{p['id']}:{got}/{rgb}")
chk(not bad_cf, "条件付き書式の式と塗り色", "; ".join(bad_cf[:3]) or f"{len(procs)} 行")

# 入力規則
dv_cells = set()
for dv in ws.data_validations.dataValidation:
    chk(dv.type == "date", "入力規則が日付型", str(dv.type))
    for rng in str(dv.sqref).split():
        from openpyxl.utils.cell import range_boundaries
        c1, r1, c2, r2 = range_boundaries(rng)
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                dv_cells.add((rr, cc))
want_dv = {(5 + i, c) for i in range(len(procs)) for c in (3, 4)}
chk(dv_cells == want_dv, "入力規則が全工程の C/D に付いている",
    f"{len(dv_cells)} セル / 期待 {len(want_dv)}")

# _data が元 CSV と等価
chk(wd.cell(row=1, column=1).value == "工程ID", "_data の 1 列目が 工程ID")
hdr = [wd.cell(row=1, column=c).value for c in range(2, len(headers) + 2)]
chk(hdr == headers, "_data が元 CSV の見出しを列順どおり保持",
    f"{len(hdr)} 列 / 元 {len(headers)} 列")
all_rows = [r for r in rows[3:] if any(c.strip() for c in r)]
bad_data = 0
for i, r in enumerate(all_rows):
    if str(wd.cell(row=2 + i, column=1).value or "") != r[H["工程ID"]]:
        bad_data += 1
        continue
    for c in range(len(headers)):
        got = wd.cell(row=2 + i, column=2 + c).value
        got = "" if got is None else str(got)
        if got != r[c]:
            bad_data += 1
            break
chk(bad_data == 0, "_data の全セルが元 CSV と等価", f"{len(all_rows)} 行 / NG {bad_data}")

out = "\n".join(lines)
print(out)
print(f"\n{'ALL PASS' if not fails else str(len(fails)) + ' FAILED'}  ({len(lines)} 項目)")
(ROOT / "out" / "inspection-openpyxl.log").write_text(out + "\n", encoding="utf-8")
sys.exit(0 if not fails else 1)
